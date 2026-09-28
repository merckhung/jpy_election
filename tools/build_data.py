"""Builds jpy_election's data files from the official sources in data/raw/.

    python tools/fetch_raw.py      # once: downloads data/raw/*
    python tools/build_data.py     # writes data/map/japan.topo.json and data/election/2026/*.json

Outputs
  data/map/japan.topo.json          TopoJSON with objects "prefectures" (47),
                                    "districts" (289 single-member districts) and
                                    "units" (municipalities / wards, clipped to
                                    districts where a municipality is split).
  data/election/2026/parties.json   parties (ja/en/zh-TW names, colours)
  data/election/2026/candidates.json   289 SMD races + 11 PR blocs
  data/election/2026/election.json  offices, PR blocs, national review, sources
  data/election/2026/results.json   official final results (MIC), unit level
  data/election/2026/results_preelection.json   empty feed (pre-election view)

Region codes
  "JP"                nation
  "13"                prefecture (JIS X 0401)
  "13-01"             single-member district (Tokyo 1st)
  "13101"             municipality/ward wholly inside one district
  "13111_03"          part of a split municipality (Ota-ku) inside district 3
"""
import json
import math
import os
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl
import shapefile
from shapely import make_valid
from shapely.geometry import MultiPolygon, Polygon, shape
from shapely.ops import unary_union
from shapely.strtree import STRtree

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
OUT_MAP = os.path.join(ROOT, "data", "map")
OUT_EL = os.path.join(ROOT, "data", "election", "2026")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jp_names  # noqa: E402  (prefecture / party tables, romanisation)

# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------

_FW_DIGITS = str.maketrans("０１２３４５６７８９", "0123456789")


def clean(s):
    """Removes ideographic variation selectors and normalises spaces."""
    if s is None:
        return ""
    s = str(s)
    s = "".join(ch for ch in s if not (0xE0100 <= ord(ch) <= 0xE01EF or 0xFE00 <= ord(ch) <= 0xFE0F))
    return s.strip()


def num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = clean(v).replace(",", "").replace(" ", "").replace("\u3000", "")
    if not s or s in ("-", "－"):
        return 0
    try:
        return int(round(float(s)))
    except ValueError:
        return 0


def ballot_name(raw):
    """'山　田　　み　き' -> ('山田', 'みき', '山田みき')."""
    s = clean(raw)
    parts = re.split(r"\u3000{2,}|\s{2,}", s)
    parts = [re.sub(r"[\s\u3000]", "", p) for p in parts if p.strip()]
    if len(parts) >= 2:
        return parts[0], "".join(parts[1:]), "".join(parts)
    one = re.sub(r"[\s\u3000]", "", s)
    return one, "", one


def squash(s):
    return re.sub(r"[\s\u3000]", "", clean(s))


# --------------------------------------------------------------------------
# MIC municipality-level workbooks
# --------------------------------------------------------------------------


def sheet_rows(ws):
    return [[clean(c) if isinstance(c, str) else c for c in row] for row in ws.iter_rows(values_only=True)]


def pref_of_sheet(rows):
    for r in rows[:8]:
        for c in r:
            if isinstance(c, str) and c.endswith("選挙管理委員会"):
                return c[: -len("選挙管理委員会")]
    return None


def leaf_rows(rows, start, ncols):
    """Returns [(name, [votes...])] for the leaf counting districts (開票区)."""
    entries = []
    for r in rows[start:]:
        c0 = r[0] if len(r) > 0 else None
        c1 = r[1] if len(r) > 1 else None
        c0 = clean(c0) if isinstance(c0, str) else ("" if c0 is None else str(c0))
        c1 = clean(c1) if isinstance(c1, str) else ("" if c1 is None else str(c1))
        name = c1 or c0
        if not name:
            continue
        if name.endswith("計") or name.startswith("※") or name.startswith("（注"):
            continue
        votes = [num(r[2 + i]) if 2 + i < len(r) else 0 for i in range(ncols)]
        entries.append({"name": name, "top": bool(c0 and not c1), "votes": votes})
    names = [e["name"] for e in entries]
    out = []
    for i, e in enumerate(entries):
        n = e["name"]
        # Aggregate rows: a designated city followed by its wards, or a ward
        # followed by its per-district parts ("北区" -> "北区第１").
        is_agg = any(m != n and m.startswith(n) for m in names)
        if e["top"] and i + 1 < len(entries) and not entries[i + 1]["top"]:
            is_agg = True
        if not is_agg:
            out.append((n, e["votes"]))
    return out


def parse_smd():
    manifest = [l.split("\t") for l in open(os.path.join(RAW, "mic_shiku_manifest.tsv"), encoding="utf-8").read().split("\n") if l.strip()]
    districts = OrderedDict()
    for kind, label, rel in manifest:
        if kind != "smd":
            continue
        wb = openpyxl.load_workbook(os.path.join(RAW, rel), read_only=True, data_only=True)
        pref_name = label
        pref = jp_names.PREF_BY_JA[pref_name]
        for ws in wb.worksheets:
            m = re.match(r"第(\d+)区", ws.title.translate(_FW_DIGITS))
            if not m:
                continue
            dno = int(m.group(1))
            rows = sheet_rows(ws)
            hdr = next(i for i, r in enumerate(rows) if r and clean(r[0]) == "候補者名")
            names = []
            for c in rows[hdr][2:]:
                if isinstance(c, str) and "得票数計" in c:
                    break
                if c is None or clean(c) == "":
                    continue
                names.append(clean(c))
            prow = next(i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].startswith("開票区名"))
            parties = [clean(c) for c in rows[prow][2 : 2 + len(names)]]
            units = leaf_rows(rows, prow + 1, len(names))
            code = f"{pref['code']}-{dno:02d}"
            districts[code] = {
                "code": code,
                "pref": pref["code"],
                "no": dno,
                "candidates": [{"raw": n, "party_ja": p} for n, p in zip(names, parties)],
                "units": units,
            }
    return districts


def parse_pr():
    manifest = [l.split("\t") for l in open(os.path.join(RAW, "mic_shiku_manifest.tsv"), encoding="utf-8").read().split("\n") if l.strip()]
    blocs = OrderedDict()
    for kind, label, rel in manifest:
        if kind != "pr":
            continue
        bloc = jp_names.BLOC_BY_JA[label]
        wb = openpyxl.load_workbook(os.path.join(RAW, rel), read_only=True, data_only=True)
        per_pref = OrderedDict()
        parties = None
        for ws in wb.worksheets:
            rows = sheet_rows(ws)
            pname = pref_of_sheet(rows)
            if not pname:
                continue
            prow = next(i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].startswith("届出番号"))
            # Party names are on the first non-empty row after the number row.
            names_row = None
            for r in rows[prow + 1 : prow + 4]:
                vals = [clean(c) for c in r[2:] if isinstance(c, str) and clean(c) and clean(c) != "党派名"]
                if len(vals) >= 3:
                    names_row = r
                    break
            ps = [clean(c) for c in names_row[2:] if isinstance(c, str) and clean(c)]
            if parties is None:
                parties = ps
            elif ps != parties:
                raise SystemExit(f"party order differs within bloc {label}: {ps} vs {parties}")
            start = next(i for i, r in enumerate(rows) if r and isinstance(r[0], str) and r[0].startswith("開票区名")) + 1
            per_pref[jp_names.PREF_BY_JA[pname]["code"]] = leaf_rows(rows, start, len(parties))
        blocs[bloc["id"]] = {"id": bloc["id"], "parties_ja": parties, "prefs": per_pref}
    return blocs


# --------------------------------------------------------------------------
# Candidate metadata from the MIC PDF (age, 新/前/元, dual candidacy)
# --------------------------------------------------------------------------

CAND_RE = re.compile(
    r"(当|落)\s+(\S.*?)\s+(\d{2,3})\s+(\S+?)\s+(新|前|元)\s+(.*?)\s*([\d,]{2,9})\s*"
    r"(?:(重)\s+([\d.]+|×))?(?=\s+(?:当|落)\s|\s*$)"
)


def parse_candidate_pdf():
    from pypdf import PdfReader

    path = os.path.join(RAW, "mic", "001061487.pdf")
    reader = PdfReader(path)
    out = []
    for page in reader.pages:
        text = page.extract_text() or ""
        for line in text.split("\n"):
            for m in CAND_RE.finditer(line.strip()):
                out.append({
                    "won": m.group(1) == "当",
                    "name": squash(m.group(2)),
                    "age": int(m.group(3)),
                    "party_ja": m.group(4),
                    "status": m.group(5),
                    "occupation": m.group(6).strip(),
                    "votes": num(m.group(7)),
                    "dual": m.group(8) == "重",
                    "sekihairitsu": None if not m.group(9) or m.group(9) == "×" else float(m.group(9)),
                    "dual_disqualified": m.group(9) == "×",
                })
    return out


# --------------------------------------------------------------------------
# Geography
# --------------------------------------------------------------------------


def decode_topojson(path):
    d = json.load(open(path, encoding="utf-8"))
    sx, sy = d["transform"]["scale"]
    tx, ty = d["transform"]["translate"]
    arcs = []
    for arc in d["arcs"]:
        x = y = 0
        pts = []
        for dx, dy in arc:
            x += dx
            y += dy
            pts.append((x * sx + tx, y * sy + ty))
        arcs.append(pts)

    def ring(idxs):
        pts = []
        for i in idxs:
            a = arcs[i] if i >= 0 else arcs[~i][::-1]
            pts.extend(a if not pts else a[1:])
        return pts

    feats = []
    for obj in d["objects"].values():
        for g in obj["geometries"]:
            polys = []
            if g["type"] == "Polygon":
                polys.append(g["arcs"])
            elif g["type"] == "MultiPolygon":
                polys.extend(g["arcs"])
            geoms = []
            for p in polys:
                rings = [ring(r) for r in p]
                if len(rings[0]) < 4:
                    continue
                geoms.append(Polygon(rings[0], [r for r in rings[1:] if len(r) >= 4]))
            if geoms:
                feats.append((g.get("properties", {}), make_valid(MultiPolygon(geoms) if len(geoms) > 1 else geoms[0])))
    return feats


def polys_only(g):
    if g.is_empty:
        return g
    if g.geom_type in ("Polygon", "MultiPolygon"):
        return g
    parts = [p for p in getattr(g, "geoms", []) if p.geom_type in ("Polygon", "MultiPolygon")]
    return unary_union(parts) if parts else Polygon()


def load_municipalities():
    feats = decode_topojson(os.path.join(RAW, "N03-21_210101.topo.json"))
    by_code = defaultdict(list)
    props = {}
    for p, g in feats:
        code = p.get("N03_007")
        if not code:
            continue
        by_code[code].append(g)
        props[code] = p
    munis = {}
    for code, gs in by_code.items():
        p = props[code]
        g = polys_only(make_valid(unary_union(gs)))
        munis[code] = {"code": code, "pref_ja": p["N03_001"], "city": p.get("N03_003") or "",
                       "name": p.get("N03_004") or "", "geom": g}
    # Hamamatsu reorganised its 7 wards into 3 on 2024-01-01 (中央区/浜名区/天竜区).
    hm = {"22131": "22138", "22132": "22138", "22133": "22138", "22134": "22138",
          "22135": "22139", "22136": "22139", "22137": "22140"}
    merged = defaultdict(list)
    for old, new in hm.items():
        if old in munis:
            merged[new].append(munis.pop(old)["geom"])
    names = {"22138": "中央区", "22139": "浜名区", "22140": "天竜区"}
    for new, gs in merged.items():
        munis[new] = {"code": new, "pref_ja": "静岡県", "city": "浜松市", "name": names[new],
                      "geom": polys_only(make_valid(unary_union(gs)))}
    return munis


def load_district_pieces():
    r = shapefile.Reader(os.path.join(RAW, "senkyoku2022", "senkyoku2022"), encoding="cp932")
    polys, codes = [], []
    for sr in r.iterShapeRecords():
        rec = sr.record
        kucode = int(rec["kucode"])
        g = shape(sr.shape.__geo_interface__)
        if g.is_empty:
            continue
        polys.append(g)
        codes.append(f"{kucode // 100:02d}-{kucode % 100:02d}")
    return polys, codes


# --------------------------------------------------------------------------
# TopoJSON writer (quantised; arcs shared between identical ring segments)
# --------------------------------------------------------------------------


class TopoWriter:
    def __init__(self, bbox, q=1_000_000):
        self.x0, self.y0, x1, y1 = bbox
        self.kx = (q - 1) / (x1 - self.x0)
        self.ky = (q - 1) / (y1 - self.y0)
        self.arcs = []
        self.arc_index = {}

    def quant(self, coords):
        out = []
        for x, y in coords:
            p = (int(round((x - self.x0) * self.kx)), int(round((y - self.y0) * self.ky)))
            if not out or out[-1] != p:
                out.append(p)
        return out

    def add_ring(self, coords):
        q = self.quant(coords)
        if len(q) < 4:
            return None
        if q[0] != q[-1]:
            q.append(q[0])
        key = tuple(q)
        rkey = tuple(reversed(q))
        if key in self.arc_index:
            return self.arc_index[key]
        if rkey in self.arc_index:
            return ~self.arc_index[rkey]
        idx = len(self.arcs)
        self.arcs.append(q)
        self.arc_index[key] = idx
        return idx

    def geometry(self, geom, props):
        polys = [geom] if geom.geom_type == "Polygon" else list(geom.geoms)
        out = []
        for p in polys:
            rings = []
            ext = self.add_ring(p.exterior.coords)
            if ext is None:
                continue
            rings.append([ext])
            for hole in p.interiors:
                h = self.add_ring(hole.coords)
                if h is not None:
                    rings.append([h])
            out.append(rings)
        if not out:
            return None
        if len(out) == 1:
            return {"type": "Polygon", "arcs": out[0], "properties": props}
        return {"type": "MultiPolygon", "arcs": out, "properties": props}

    def to_json(self, objects):
        arcs = []
        for a in self.arcs:
            enc = [list(a[0])]
            for i in range(1, len(a)):
                enc.append([a[i][0] - a[i - 1][0], a[i][1] - a[i - 1][1]])
            arcs.append(enc)
        return {
            "type": "Topology",
            "transform": {"scale": [1 / self.kx, 1 / self.ky], "translate": [self.x0, self.y0]},
            "objects": {k: {"type": "GeometryCollection", "geometries": v} for k, v in objects.items()},
            "arcs": arcs,
        }


def round_geom(g, tol):
    g = g.simplify(tol, preserve_topology=True) if tol > 0 else g
    return polys_only(make_valid(g))


# --------------------------------------------------------------------------
# Main build
# --------------------------------------------------------------------------


def main():
    os.makedirs(OUT_MAP, exist_ok=True)
    os.makedirs(OUT_EL, exist_ok=True)
    print("parsing MIC SMD workbooks ...")
    districts = parse_smd()
    print(f"  {len(districts)} districts, {sum(len(d['candidates']) for d in districts.values())} candidates, "
          f"{sum(len(d['units']) for d in districts.values())} counting districts")
    print("parsing MIC PR workbooks ...")
    blocs = parse_pr()
    print(f"  {len(blocs)} blocs")
    print("parsing candidate PDF ...")
    pdf = parse_candidate_pdf()
    print(f"  {len(pdf)} candidate rows")

    # ---- municipalities and matching of counting districts
    print("loading municipalities ...")
    munis = load_municipalities()
    name_index = defaultdict(dict)  # pref code -> name -> muni code
    for code, m in munis.items():
        pc = code[:2]
        keys = set()
        if m["city"] and m["city"].endswith("市"):
            keys.add(m["city"] + m["name"])  # 札幌市中央区
        keys.add(m["name"])
        if m["city"]:
            keys.add(m["city"] + m["name"])  # 神石郡神石高原町
        for k in keys:
            name_index[pc].setdefault(k, code)

    def match_unit(pref_code, name):
        base = re.sub(r"第[0-9０-９]+$", "", name)
        pref_ja = jp_names.PREF_BY_CODE[pref_code]["ja"]
        if base.endswith(pref_ja) and len(base) > len(pref_ja):
            base = base[: -len(pref_ja)]
        base = base.replace("ヶ", "ケ") if base not in name_index[pref_code] else base
        for cand in (name, base, base.replace("ケ", "ヶ")):
            if cand in name_index[pref_code]:
                return name_index[pref_code][cand]
        return jp_names.UNIT_ALIASES.get((pref_code, base))

    unmatched = []
    muni_districts = defaultdict(set)
    for d in districts.values():
        for i, (uname, votes) in enumerate(d["units"]):
            mc = match_unit(d["pref"], uname)
            if not mc:
                unmatched.append((d["code"], uname))
            else:
                muni_districts[mc].add(d["code"])
    if unmatched:
        print("UNMATCHED SMD counting districts:")
        for u in unmatched:
            print("  ", u)
        raise SystemExit(1)

    for d in districts.values():
        agg = OrderedDict()
        for uname, votes in d["units"]:
            mc = match_unit(d["pref"], uname)
            key = mc if len(muni_districts[mc]) == 1 else f"{mc}_{d['no']:02d}"
            if key not in agg:
                agg[key] = {"muni": mc, "names": [], "votes": [0] * len(votes)}
            agg[key]["names"].append(uname)
            agg[key]["votes"] = [a + b for a, b in zip(agg[key]["votes"], votes)]
        d["unit_votes"] = agg

    split = {mc: sorted(ds) for mc, ds in muni_districts.items() if len(ds) > 1}
    print(f"  {len(muni_districts)} municipalities with votes, {len(split)} split across districts")

    if "--election-only" in sys.argv:
        import build_election  # noqa: E402

        build_election.write_all(districts=districts, blocs=blocs, pdf=pdf, munis=munis,
                                 match_unit=match_unit, muni_districts=muni_districts, unit_names={},
                                 out_dir=OUT_EL, raw=RAW)
        return

    # ---- geometry of units
    print("loading district polygons (this takes a while) ...")
    dpolys, dcodes = load_district_pieces()
    tree = STRtree(dpolys)
    print(f"  {len(dpolys)} district pieces")

    unit_geom = {}
    for mc, ds in muni_districts.items():
        g = munis[mc]["geom"]
        if mc not in split:
            (dc,) = ds
            unit_geom[mc] = (dc, g)
            continue
        idx = tree.query(g)
        parts = {}
        for dc in split[mc]:
            cand = [dpolys[i] for i in idx if dcodes[i] == dc]
            if not cand:
                continue
            dg = make_valid(unary_union(cand)).buffer(0)
            parts[dc] = polys_only(g.intersection(dg))
        # Clean overlaps and give leftovers (coastline mismatch) to the biggest part.
        order = sorted(parts, key=lambda k: -parts[k].area)
        used = Polygon()
        for k in order:
            parts[k] = polys_only(parts[k].difference(used))
            used = used.union(parts[k])
        left = polys_only(g.difference(used))
        if order and not left.is_empty:
            parts[order[0]] = polys_only(parts[order[0]].union(left))
        for dc in split[mc]:
            no = int(dc.split("-")[1])
            p = parts.get(dc)
            if p is None or p.is_empty:
                print(f"  warning: empty part {mc} in {dc}; using a point buffer")
                p = g.representative_point().buffer(0.002)
            unit_geom[f"{mc}_{no:02d}"] = (dc, p)

    # Municipalities without votes (the Northern Territories) join the
    # district that contains them, so the map has no holes.
    for mc, m in munis.items():
        if mc in muni_districts:
            continue
        idx = tree.query(m["geom"].representative_point())
        dc = dcodes[idx[0]] if len(idx) else None
        if dc is None:
            # Nearest district piece.
            dc = dcodes[tree.nearest(m["geom"].representative_point())]
        unit_geom[mc] = (dc, m["geom"])
        print(f"  no-vote unit {mc} {m['city']}{m['name']} -> {dc}")

    print("building TopoJSON ...")
    # Simplify units lightly (split parts carry the detailed district lines).
    ug = {k: round_geom(g, 0.0004 if "_" in k else 0.0) for k, (dc, g) in unit_geom.items()}
    dist_geom = defaultdict(list)
    for k, (dc, _) in unit_geom.items():
        dist_geom[dc].append(ug[k])
    dist_geom = {dc: polys_only(make_valid(unary_union(gs))) for dc, gs in dist_geom.items()}
    pref_geom = defaultdict(list)
    for dc, g in dist_geom.items():
        pref_geom[dc[:2]].append(g)
    pref_geom = {pc: polys_only(make_valid(unary_union(gs))) for pc, gs in pref_geom.items()}

    allg = unary_union(list(pref_geom.values()))
    tw = TopoWriter(allg.bounds)
    objs = {"prefectures": [], "districts": [], "units": []}
    for pc in sorted(pref_geom):
        p = jp_names.PREF_BY_CODE[pc]
        objs["prefectures"].append(tw.geometry(pref_geom[pc], {"code": pc, "name_ja": p["ja"], "name_en": p["en"],
                                                               "name_zh": jp_names.to_zh(p["ja"])}))
    for dc in sorted(dist_geom):
        pc, no = dc.split("-")
        p = jp_names.PREF_BY_CODE[pc]
        ja = f"{p['short']}{int(no)}区"
        objs["districts"].append(tw.geometry(dist_geom[dc], {"code": dc, "pref": pc, "name_ja": ja,
                                                             "name_en": f"{p['en']} {int(no)}",
                                                             "name_zh": jp_names.to_zh(ja)}))
    unit_names = {}
    for k in sorted(unit_geom):
        dc = unit_geom[k][0]
        mc = k.split("_")[0]
        m = munis[mc]
        ja = (m["city"] + m["name"]) if m["city"].endswith("市") else m["name"]
        en = jp_names.muni_en(mc, ja)
        if "_" in k:
            ja_label = f"{ja}（{int(k.split('_')[1])}区）"
            en_label = f"{en} (part, {int(k.split('_')[1])})"
        else:
            ja_label, en_label = ja, en
        unit_names[k] = ja_label
        g = tw.geometry(ug[k], {"code": k, "district": dc, "pref": dc[:2], "muni": mc, "name_ja": ja_label,
                                "name_en": en_label, "name_zh": jp_names.to_zh(ja_label)})
        if g:
            objs["units"].append(g)
    topo = tw.to_json(objs)
    with open(os.path.join(OUT_MAP, "japan.topo.json"), "w", encoding="utf-8") as f:
        json.dump(topo, f, ensure_ascii=False, separators=(",", ":"))
    print(f"  wrote japan.topo.json: {len(objs['prefectures'])} prefectures, {len(objs['districts'])} districts, "
          f"{len(objs['units'])} units, {len(topo['arcs'])} arcs, "
          f"{os.path.getsize(os.path.join(OUT_MAP, 'japan.topo.json')) / 1e6:.1f} MB")

    # ---- election JSON
    import build_election  # noqa: E402

    build_election.write_all(districts=districts, blocs=blocs, pdf=pdf, munis=munis,
                             match_unit=match_unit, muni_districts=muni_districts, unit_names=unit_names,
                             out_dir=OUT_EL, raw=RAW)


if __name__ == "__main__":
    main()
