"""Writes data/election/2026/*.json for jpy_election (called by build_data.py).

All numbers come from the official MIC (総務省) result tables of the 51st
House of Representatives general election and the 27th Supreme Court
national review, 2026-02-08. Two quantities are *not* published at the
granularity the app wants and are estimated (documented in the JSON):

  * eligible voters per single-member district: prefecture total (MIC) x the
    district's share of the prefecture's Japanese population (2020 census,
    as compiled for the 2022 redistricting);
  * eligible voters / ballots cast per counting unit: the district figure
    split in proportion to the unit's valid votes (uniform turnout inside a
    district).

Vote counts per candidate, party and counting unit are exact (MIC rounds
fractional "按分" votes; we round each unit to an integer).
"""
import json
import os
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jp_names  # noqa: E402

ELECTION = OrderedDict(
    id="2026-shugiin51",
    date="2026-02-08",
    polls_open="07:00",
    polls_close="20:00",
    timezone="Asia/Tokyo",
    name_ja="第51回衆議院議員総選挙",
    name_en="2026 Japanese general election (51st House of Representatives)",
    name_zh="第51屆日本眾議院議員總選舉",
)

_FW = str.maketrans("０１２３４５６７８９", "0123456789")


def clean(s):
    if s is None:
        return ""
    s = str(s)
    s = "".join(ch for ch in s if not (0xE0100 <= ord(ch) <= 0xE01EF or 0xFE00 <= ord(ch) <= 0xFE0F))
    return s.strip()


def squash(s):
    s = re.sub(r"[\s\u3000()（）]", "", clean(s))
    return s.replace("ッ", "ツ").replace("ヶ", "ケ")


def num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = clean(v).replace(",", "").replace(" ", "").replace("\u3000", "")
    try:
        return int(round(float(s)))
    except ValueError:
        return 0


def ballot_name(raw):
    s = clean(raw)
    parts = re.split(r"\u3000{2,}|\s{2,}", s)
    parts = [re.sub(r"[\s\u3000]", "", p) for p in parts if p.strip()]
    if len(parts) >= 2:
        return parts[0], "".join(parts[1:])
    return re.sub(r"[\s\u3000]", "", s), ""


def is_kana(s):
    return bool(s) and all(("\u3040" <= c <= "\u30ff") or c in "ー・" for c in s)


_KAKASI = None


def kanji_reading(s):
    global _KAKASI
    if _KAKASI is None:
        import pykakasi

        _KAKASI = pykakasi.kakasi()
    return "".join(x["hira"] for x in _KAKASI.convert(s))


def ordinal(n):
    if 10 <= n % 100 <= 20:
        return f"{n}th"
    return f"{n}{ {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th') }"


def party_code(ja):
    p = jp_names.PARTY_BY_JA.get(ja)
    return p[0] if p else "MINOR"


def largest_remainder(total, weights):
    """Splits integer `total` in proportion to `weights` (exact sum)."""
    s = sum(weights)
    if s <= 0:
        return [0] * len(weights)
    raw = [total * w / s for w in weights]
    out = [int(x) for x in raw]
    rest = total - sum(out)
    order = sorted(range(len(raw)), key=lambda i: -(raw[i] - out[i]))
    for i in order[:rest]:
        out[i] += 1
    return out


# --------------------------------------------------------------------------
# PDF helpers (pdfplumber word boxes)
# --------------------------------------------------------------------------


def page_lines(page, tol=3):
    words = page.extract_words(x_tolerance=1.5, y_tolerance=2)
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0] - w["top"]) <= tol:
            lines[-1][1].append(w)
        else:
            lines.append([w["top"], [w]])
    return [(top, sorted(ws, key=lambda w: w["x0"])) for top, ws in lines]


def split_by_gap(words):
    """Splits name words into (surname, given) at the widest horizontal gap."""
    if len(words) <= 1:
        return words, []
    gaps = [(words[i + 1]["x0"] - words[i]["x1"], i) for i in range(len(words) - 1)]
    _, i = max(gaps)
    return words[: i + 1], words[i + 1 :]


VOTES_RE = re.compile(r"^\d{1,3}(,\d{3})*(\.\d+)?$")


def parse_candidates_pdf(raw):
    """Candidate rows from the MIC candidate PDF (001061487), column-aware.

    Returns a list of dicts: name (ballot name, squashed), votes, won, age,
    party_ja, status (新/前/元), occupation, dual, dual_disqualified,
    sekihairitsu, read_sur / read_given (furigana) and legal (戸籍 name).
    """
    import pdfplumber

    out = []
    with pdfplumber.open(os.path.join(raw, "mic", "001061487.pdf")) as pdf:
        for page in pdf.pages:
            lines = page_lines(page)
            mid = page.width / 2
            for col in (0, 1):
                lo, hi = (0, mid - 20) if col == 0 else (mid - 20, page.width)
                cl = [(t, [w for w in ws if lo <= w["x0"] < hi]) for t, ws in lines]
                cl = [(t, ws) for t, ws in cl if ws]
                for li, (top, ws) in enumerate(cl):
                    if ws[0]["text"] not in ("当", "落"):
                        continue
                    x0 = ws[0]["x0"]
                    rel = lambda w: w["x0"] - x0  # noqa: E731
                    ages = [w for w in ws[1:] if re.fullmatch(r"\d{2,3}", w["text"]) and 100 < rel(w) < 160]
                    age_x = rel(ages[0]) if ages else 145
                    name_ws = [w for w in ws[1:] if 10 <= rel(w) < age_x - 2]
                    party_ws = [w for w in ws if age_x + 10 <= rel(w) < 270]
                    status_ws = [w for w in ws if 270 <= rel(w) < 288 and w["text"] in ("新", "前", "元")]
                    votes = [w for w in ws if VOTES_RE.match(w["text"]) and 340 <= rel(w) < 420]
                    dual_ws = [w for w in ws if 415 <= rel(w) < 440 and w["text"] == "重"]
                    seki_ws = [w for w in ws if rel(w) >= 438 and re.fullmatch(r"[\d.]+|×", w["text"])]
                    if not name_ws or not votes:
                        continue
                    occ = lambda wl: "".join(w["text"] for w in wl if 286 <= rel(w) < 360)  # noqa: E731
                    sur_ws, giv_ws = split_by_gap(name_ws)
                    furi, legal, occ_above, occ_below = [], "", "", ""
                    for t2, ws2 in cl[max(0, li - 3) : li + 3]:
                        in_name = [w for w in ws2 if 10 <= rel(w) < age_x]
                        if top - 15 <= t2 <= top - 6:
                            furi += [w for w in in_name if is_kana(w["text"])]
                            occ_above = occ(ws2)
                        if top + 6 <= t2 <= top + 15:
                            if in_name and in_name[0]["text"].startswith("("):
                                legal = re.sub(r"[()（）]", "", "".join(w["text"] for w in in_name))
                            occ_below = occ(ws2)
                    rs = rg = ""
                    for f in furi:
                        starts = [(abs(f["x0"] - p[0]["x0"]), k) for k, p in enumerate((sur_ws, giv_ws)) if p]
                        if min(starts)[1] == 0:
                            rs += f["text"]
                        else:
                            rg += f["text"]
                    seki = seki_ws[0]["text"] if seki_ws else ""
                    out.append({
                        "name": squash("".join(w["text"] for w in name_ws)),
                        "votes": num(votes[0]["text"]),
                        "won": ws[0]["text"] == "当",
                        "age": int(ages[0]["text"]) if ages else None,
                        "party_ja": "".join(w["text"] for w in party_ws),
                        "status": status_ws[0]["text"] if status_ws else "",
                        "occupation": occ_above + occ(ws) + occ_below,
                        "dual": bool(dual_ws),
                        "dual_disqualified": seki == "×",
                        "sekihairitsu": float(seki) if seki and seki != "×" else None,
                        "read_sur": rs,
                        "read_given": rg,
                        "legal": legal,
                    })
    return out


def parse_pr_lists(raw):
    """Party lists per bloc from the MIC PR result PDF (001061485).

    Returns {bloc_id: {party_ja: {"seats": n, "list": [entry...]}}} where an
    entry is {"no", "name", "surname", "given", "smd": "当"/"落"/"", "seki",
    "order"} and "order" is the election order for PR winners (else None).
    """
    import pdfplumber

    blocs = OrderedDict()
    label_to_id = {b["label"]: b["id"] for b in jp_names.BLOC_BY_JA.values()}
    with pdfplumber.open(os.path.join(raw, "mic", "001061485.pdf")) as pdf:
        for page in pdf.pages:
            lines = page_lines(page)
            bloc = None
            for _, ws in lines[:6]:
                for w in ws:
                    if w["text"] in label_to_id:
                        bloc = label_to_id[w["text"]]
            if not bloc:
                continue
            hdr = next((ws for _, ws in lines if any(w["text"].startswith("名") and len(w["text"]) > 1 for w in ws)
                        and any(w["text"] == "党" for w in ws)), None)
            if not hdr:
                continue
            cols = [w["x0"] for w in hdr if w["text"] == "党"]
            names = [w["text"][1:] for w in hdr if w["text"].startswith("名") and len(w["text"]) > 1]
            width = (cols[1] - cols[0]) if len(cols) > 1 else 190
            seats_line = next(ws for _, ws in lines if sum(1 for w in ws if w["text"] == "選") >= 1
                              and any(w["text"] == "当" for w in ws) and any(w["text"] == "人" for w in ws))
            header_top = next(t for t, ws in lines if any(w["text"] == "順位" for w in ws))
            for ci, (cx, pname) in enumerate(zip(cols, names)):
                in_col = lambda w: cx - 6 <= w["x0"] < cx + width - 6  # noqa: E731
                sw = [w for w in seats_line if in_col(w)]
                seats = 0
                for k, w in enumerate(sw):
                    if w["text"] == "人" and k > 0 and sw[k - 1]["text"].isdigit():
                        seats = int(sw[k - 1]["text"])
                        break
                entry_list = []
                for top, ws in lines:
                    if top <= header_top:
                        continue
                    cw = [w for w in ws if in_col(w)]
                    if not cw or not cw[0]["text"].isdigit() or cw[0]["x0"] - cx > 14:
                        continue
                    no = int(cw[0]["text"])
                    rel = lambda w: w["x0"] - cx  # noqa: E731
                    name_ws = [w for w in cw[1:] if 12 <= rel(w) < 110]
                    order_ws = [w for w in cw if 110 <= rel(w) < 128 and w["text"].isdigit()]
                    smd_ws = [w for w in cw if 128 <= rel(w) < 148 and w["text"] in ("当", "落")]
                    seki_ws = [w for w in cw if rel(w) >= 148]
                    sur, giv = split_by_gap(name_ws)
                    seki = seki_ws[0]["text"] if seki_ws else ""
                    entry_list.append({
                        "no": no,
                        "name": "".join(w["text"] for w in name_ws),
                        "surname": "".join(w["text"] for w in sur),
                        "given": "".join(w["text"] for w in giv),
                        "order": int(order_ws[0]["text"]) if order_ws else None,
                        "smd": smd_ws[0]["text"] if smd_ws else "",
                        "seki": None if not seki or seki == "×" else float(seki),
                        "disqualified": seki == "×",
                    })
                b = blocs.setdefault(bloc, OrderedDict())
                if pname in b:
                    b[pname]["list"] += entry_list
                else:
                    b[pname] = {"seats": seats, "list": entry_list}
    return blocs


# --------------------------------------------------------------------------
# Prefecture statistics
# --------------------------------------------------------------------------


def xl_rows(path, sheet=0):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def pref_table(path, first_num_col_after_label=True):
    """{pref code: [numbers...]} for MIC per-prefecture tables."""
    out = {}
    for r in xl_rows(path):
        for k in range(min(3, len(r))):
            c = clean(r[k]) if isinstance(r[k], str) else ""
            c = c.replace("\u3000", "").replace(" ", "")
            if c in jp_names.PREF_BY_JA:
                out[jp_names.PREF_BY_JA[c]["code"]] = [num(x) for x in r[k + 1 :]]
                break
    return out


def district_population(raw):
    out = {}
    for r in xl_rows(os.path.join(raw, "senkyoku_ichiran.xlsx"))[4:]:
        if not r or not isinstance(r[0], (int, float)) or not isinstance(r[1], (int, float)):
            continue
        out[f"{int(r[0]):02d}-{int(r[1]):02d}"] = num(r[4])
    return out


def incumbents_2024(raw):
    """2024 (50th) winners per district: {code: (name, party_ja)}."""
    out = {}
    path = os.path.join(raw, "mic50_shiku_manifest.tsv")
    if not os.path.exists(path):
        return out
    for line in open(path, encoding="utf-8").read().split("\n"):
        if not line.strip():
            continue
        _, label, rel = line.split("\t")
        pref = jp_names.PREF_BY_JA.get(label)
        if not pref:
            continue
        wb = openpyxl.load_workbook(os.path.join(raw, rel), read_only=True, data_only=True)
        for ws in wb.worksheets:
            m = re.search(r"第(\d+)区", ws.title.translate(_FW))
            if not m:
                continue
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            hi = next((i for i, r in enumerate(rows) if r and clean(r[0]) == "候補者名"), None)
            if hi is None:
                continue
            names = []
            for c in rows[hi][1:]:
                if isinstance(c, str) and "得票数計" in c:
                    break
                names.append(clean(c))
            parties = [clean(c) for c in rows[hi + 1][1 : 1 + len(names)]]
            tot = next((r for r in rows if r and isinstance(r[0], str) and clean(r[0]).endswith("合計")), None)
            if not tot:
                continue
            votes = [num(x) for x in tot[1 : 1 + len(names)]]
            w = max(range(len(names)), key=lambda i: votes[i])
            out[f"{pref['code']}-{int(m.group(1)):02d}"] = (names[w], parties[w])
    return out


# --------------------------------------------------------------------------
# Main entry
# --------------------------------------------------------------------------


def romanize(read_sur, read_given):
    s = jp_names.cap(jp_names.romaji(read_sur)) if read_sur else ""
    g = jp_names.cap(jp_names.romaji(read_given)) if read_given else ""
    return (g + " " + s).strip()


def write_json(path, obj):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write("\n")


def write_all(districts, blocs, pdf, munis, match_unit, muni_districts, unit_names, out_dir, raw, **_):
    os.makedirs(out_dir, exist_ok=True)
    print("build_election: candidate metadata / PR lists from PDFs ...")
    pdf = parse_candidates_pdf(raw)  # replaces the regex-based rows from build_data
    pr_lists = parse_pr_lists(raw)
    inc24 = incumbents_2024(raw)
    print(f"  {len(pdf)} candidate rows, {sum(len(v) for v in pr_lists.values())} party lists, "
          f"{len(inc24)} 2024 winners")

    # ---------------- parties
    used_codes = set()
    for d in districts.values():
        for c in d["candidates"]:
            used_codes.add(party_code(c["party_ja"]))
    for b in blocs.values():
        for p in b["parties_ja"]:
            used_codes.add(party_code(p))
    parties = []
    for code, ja, sja, en, sen, zh, szh, color in jp_names.PARTIES:
        if code not in used_codes:
            continue
        parties.append(OrderedDict(code=code, name_ja=ja, short_ja=sja, name_en=en, short_en=sen,
                                   name_zh=zh, short_zh=szh, color=color))
    write_json(os.path.join(out_dir, "parties.json"),
               OrderedDict(schema="jpy_election.parties/v1", parties=parties))

    # ---------------- PDF metadata index
    pdf_by_name = defaultdict(list)
    for row in pdf:
        pdf_by_name[squash(row["name"])].append(row)

    # ---------------- SMD candidates
    smd_races = []
    cand_index = {}  # (bloc, party, squashed name) -> candidate dict (for PR revival)
    guesses = 0
    for code, d in districts.items():
        pref = jp_names.PREF_BY_CODE[d["pref"]]
        bloc = next(b for b in jp_names.BLOC_BY_ID.values() if d["pref"] in b["prefs"])
        totals = [0] * len(d["candidates"])
        for u in d["unit_votes"].values():
            totals = [a + b for a, b in zip(totals, u["votes"])]
        cands = []
        for i, c in enumerate(d["candidates"]):
            sur, giv = ballot_name(c["raw"])
            full = sur + giv
            rows = pdf_by_name.get(squash(full), [])
            meta = min(rows, key=lambda r: abs(r["votes"] - totals[i])) if rows else None
            if meta is None or abs(meta["votes"] - totals[i]) > 50:
                print(f"  warning: no PDF metadata for {code} {full} ({totals[i]})")
                meta = {}
            rd = meta
            rs = rd.get("read_sur") or (sur if is_kana(sur) else "")
            rg = rd.get("read_given") or (giv if is_kana(giv) else "")
            guess = False
            if not rs and sur:
                rs, guess = kanji_reading(sur), True
            if not rg and giv:
                rg, guess = kanji_reading(giv), True
            guesses += guess
            legal = rd.get("legal") or ""
            pc = party_code(c["party_ja"])
            cand = OrderedDict(
                id=f"{code}-{i + 1:02d}",
                name_ja=full,
                surname_ja=sur,
                given_ja=giv,
                name_legal=legal or full,
                kana=f"{jp_names.hira_to_kata(rs)} {jp_names.hira_to_kata(rg)}".strip(),
                name_en=romanize(rs, rg),
                name_zh=jp_names.to_zh(legal or full),
                romanization_guess=guess,
                party=pc,
                party_ja=c["party_ja"],
                age=meta.get("age"),
                status=meta.get("status", ""),
                occupation=meta.get("occupation", ""),
                dual=bool(meta.get("dual")),
                dual_disqualified=bool(meta.get("dual_disqualified")),
                sekihairitsu=meta.get("sekihairitsu"),
                incumbent=meta.get("status") == "前",
                pr_elected=False,
                photo=f"assets/photos/{code}-{i + 1:02d}.jpg",
                note="",
            )
            cands.append(cand)
            cand_index.setdefault((bloc["id"], pc, squash(full)), []).append(cand)
        inc = inc24.get(code)
        incumbent = None
        if inc:
            iname, iparty = inc
            ipc = jp_names.PARTY_2024.get(iparty, party_code(iparty))
            match = [c for c in cands if squash(c["name_ja"]) == squash(iname)]
            note = ""
            if match:
                ipc = match[0]["party"]
            if iparty in ("立憲民主党", "公明党"):
                note = f"2024: {iparty}"
            incumbent = OrderedDict(name_ja=squash(iname), party=ipc, running=bool(match), note=note)
        ja = f"{pref['short']}{d['no']}区"
        smd_races.append(OrderedDict(
            id=code, kind="smd", region=code, pref=d["pref"], bloc=bloc["id"], seats=1,
            name_ja=ja, name_en=f"{pref['en']} {ordinal(d['no'])}", name_zh=jp_names.to_zh(ja),
            incumbent=incumbent, candidates=cands,
        ))
    print(f"  {sum(len(r['candidates']) for r in smd_races)} SMD candidates ({guesses} romanisation guesses)")

    # ---------------- PR races
    pr_races = []
    for bid, b in blocs.items():
        bl = jp_names.BLOC_BY_ID[bid]
        lists = pr_lists.get(bid, {})
        cands = []
        for pja in b["parties_ja"]:
            pc = party_code(pja)
            p = jp_names.PARTY_BY_CODE[pc]
            plist = []
            for e in lists.get(pja, {"list": []})["list"]:
                ent = OrderedDict(no=e["no"], name_ja=e["name"], surname_ja=e["surname"], given_ja=e["given"],
                                  smd=None, order=e["order"], disqualified=e["disqualified"])
                hits = cand_index.get((bid, pc, squash(e["name"])), [])
                if len(hits) != 1 and e["seki"] is not None:
                    hits = [c for (b2, p2, _), cl in cand_index.items() if b2 == bid and p2 == pc for c in cl
                            if c["sekihairitsu"] is not None and abs(c["sekihairitsu"] - e["seki"]) < 0.0005]
                if len(hits) == 1:
                    ent["smd"] = hits[0]["id"].rsplit("-", 1)[0]
                    ent["candidate"] = hits[0]["id"]
                    if e["order"] is not None:
                        hits[0]["pr_elected"] = True
                elif e["smd"]:
                    print(f"  warning: dual candidate not matched: {bid} {pja} {e['name']} ({len(hits)})")
                plist.append(ent)
            cands.append(OrderedDict(
                id=f"pr-{bid}-{pc}", name_ja=pja, name_en=p[3], name_zh=p[5], party=pc,
                seats_won=lists.get(pja, {}).get("seats"), list=plist,
            ))
        pr_races.append(OrderedDict(
            id=f"pr-{bid}", kind="pr", region="JP", bloc=bid, prefs=bl["prefs"], seats=bl["seats"],
            name_ja=f"比例{bl['ja']}", name_en=f"PR {bl['en']} bloc", name_zh=f"比例{bl['zh']}區",
            candidates=cands,
        ))
        won = sum((c["seats_won"] or 0) for c in cands)
        elected = sum(1 for c in cands for e in c["list"] if e["order"] is not None)
        if won != bl["seats"] or elected != bl["seats"]:
            print(f"  warning: bloc {bid}: seats {bl['seats']}, parties won {won}, elected names {elected}")

    # Validate against MIC national totals (winners by party).
    smd_w = defaultdict(int)
    for r in smd_races:
        tot = defaultdict(int)
        for u in districts[r["id"]]["unit_votes"].values():
            for i, v in enumerate(u["votes"]):
                tot[i] += v
        w = max(tot, key=lambda i: tot[i])
        smd_w[r["candidates"][w]["party"]] += 1
    pr_w = defaultdict(int)
    for r in pr_races:
        for c in r["candidates"]:
            pr_w[c["party"]] += c["seats_won"] or 0
    print("  SMD winners:", dict(sorted(smd_w.items(), key=lambda kv: -kv[1])))
    print("  PR seats:   ", dict(sorted(pr_w.items(), key=lambda kv: -kv[1])))

    write_json(os.path.join(out_dir, "candidates.json"), OrderedDict(
        schema="jpy_election.candidates/v1",
        election=ELECTION,
        status_note=("Official candidates and results of the 51st general election (MIC, 2026-02-13). "
                     "English names are romanised from the official furigana; entries with "
                     "romanization_guess=true were romanised from kanji automatically."),
        races=smd_races + pr_races,
    ))

    # ---------------- results
    stats_smd = pref_table(os.path.join(raw, "mic", "001061471.xlsx"))  # 男 女 計 | 投票者 男 女 計 | ...
    stats_pr = pref_table(os.path.join(raw, "mic", "001061473.xlsx"))
    stats_rev = pref_table(os.path.join(raw, "mic", "001061488.xlsx"))
    rev_votes = pref_table(os.path.join(raw, "mic", "001061490.xlsx"))
    dpop = district_population(raw)

    pref_valid_smd = defaultdict(int)
    dist_valid = {}
    for code, d in districts.items():
        dist_valid[code] = sum(sum(u["votes"]) for u in d["unit_votes"].values())
        pref_valid_smd[d["pref"]] += dist_valid[code]
    pref_pop = defaultdict(int)
    for code in districts:
        pref_pop[code[:2]] += dpop.get(code, 0)

    # Eligible voters per district (estimate) with exact prefecture totals.
    dist_elig = {}
    for pc in sorted(set(c[:2] for c in districts)):
        codes = sorted(c for c in districts if c[:2] == pc)
        elig = stats_smd[pc][2]
        for c, e in zip(codes, largest_remainder(elig, [dpop.get(c, 1) for c in codes])):
            dist_elig[c] = e

    unit_elig = {}
    races_json = OrderedDict()
    for code, d in districts.items():
        pc = d["pref"]
        cands = next(r for r in smd_races if r["id"] == code)["candidates"]
        keys = list(d["unit_votes"].keys())
        valid = [sum(d["unit_votes"][k]["votes"]) for k in keys]
        ball_total = round(dist_valid[code] * stats_smd[pc][5] / max(1, pref_valid_smd[pc]))
        eligs = largest_remainder(dist_elig[code], valid)
        balls = largest_remainder(ball_total, valid)
        regions = OrderedDict()
        for k, e, bcast in zip(keys, eligs, balls):
            unit_elig[k] = e
            regions[k] = OrderedDict(
                votes=OrderedDict((c["id"], v) for c, v in zip(cands, d["unit_votes"][k]["votes"])),
                eligible=e, ballots_cast=max(bcast, sum(d["unit_votes"][k]["votes"])),
                units_counted=1, units_total=1)
        races_json[code] = OrderedDict(regions=regions)

    # PR: map counting-district names (incl. "第N" parts) to unit keys.
    part_district = {}
    for code, d in districts.items():
        for k, u in d["unit_votes"].items():
            for n in u["names"]:
                part_district[(d["pref"], n)] = k
    for bid, b in blocs.items():
        race = next(r for r in pr_races if r["bloc"] == bid)
        ids = [c["id"] for c in race["candidates"]]
        unit_votes = OrderedDict()
        for pc, rows in b["prefs"].items():
            for uname, votes in rows:
                key = part_district.get((pc, uname))
                if key is None and sum(votes) == 0:
                    continue  # empty 郡 subtotal rows
                if key is None:
                    mc = match_unit(pc, uname)
                    if mc is None:
                        raise SystemExit(f"PR unit not matched: {bid} {pc} {uname}")
                    key = mc
                    if len(muni_districts.get(mc, ())) > 1:
                        print(f"  warning: PR row {uname} of split municipality {mc} has no part suffix")
                        key = f"{mc}_{int(sorted(muni_districts[mc])[0].split('-')[1]):02d}"
                acc = unit_votes.setdefault(key, [0] * len(ids))
                for i, v in enumerate(votes):
                    acc[i] += v
        pref_valid_pr = defaultdict(int)
        for k, v in unit_votes.items():
            pref_valid_pr[k[:2]] += sum(v)
        regions = OrderedDict()
        for k, v in unit_votes.items():
            pc = k[:2]
            bcast = round(sum(v) * stats_pr[pc][6] / max(1, pref_valid_pr[pc]))
            regions[k] = OrderedDict(votes=OrderedDict(zip(ids, v)), eligible=unit_elig.get(k, 0),
                                     ballots_cast=max(bcast, sum(v)), units_counted=1, units_total=1)
        races_json[race["id"]] = OrderedDict(regions=regions)

    # National review, per prefecture.
    units_per_pref = defaultdict(int)
    for k in unit_elig:
        units_per_pref[k[:2]] += 1
    refs_json = OrderedDict()
    referendums = []
    for j, (jid, ja, en, zh, bio_ja, bio_en, bio_zh) in enumerate(jp_names.REVIEW_JUSTICES):
        rid = f"review-{jid}"
        regions = OrderedDict()
        for pc in sorted(rev_votes):
            row = rev_votes[pc]
            regions[pc] = OrderedDict(agree=row[j * 4], disagree=row[j * 4 + 1], eligible=stats_rev[pc][2],
                                      ballots_cast=stats_rev[pc][5], units_counted=units_per_pref[pc],
                                      units_total=units_per_pref[pc])
        refs_json[rid] = OrderedDict(regions=regions)
        referendums.append(OrderedDict(
            id=rid, kind="review", scope="national", justice_ja=ja, justice_en=en, justice_zh=zh,
            question_ja=f"最高裁判所裁判官 {ja} を罷免することを可とするか",
            question_en=f"Should Supreme Court Justice {en} be dismissed?",
            question_zh=f"是否罷免最高法院法官 {zh}？",
            bio_ja=bio_ja, bio_en=bio_en, bio_zh=bio_zh,
            note=("A justice is dismissed only if 'dismiss' marks outnumber blank ballots "
                  "(and turnout is at least 1% of the electorate). English/Chinese texts are unofficial."),
        ))
        agree = sum(r["agree"] for r in regions.values())
        dis = sum(r["disagree"] for r in regions.values())
        print(f"  review {ja}: dismiss {agree:,} / keep {dis:,} ({100 * agree / (agree + dis):.2f}%)")

    write_json(os.path.join(out_dir, "results.json"), OrderedDict(
        schema="jpy_election.results/v1",
        status="final",
        source=("MIC (総務省) official results, 51st general election & 27th national review, 2026-02-08. "
                "Vote counts are exact per counting unit; eligible voters and ballots cast below the "
                "prefecture level are estimates (see tools/build_election.py)."),
        updated_at="2026-02-13T00:00:00+09:00",
        races=races_json,
        referendums=refs_json,
        declarations=[],
    ))
    write_json(os.path.join(out_dir, "results_preelection.json"), OrderedDict(
        schema="jpy_election.results/v1",
        status="pre-election",
        source=("No votes yet: polls open 2026-02-08 07:00-20:00 (Asia/Tokyo). Replace this file "
                "(or pass --results) with a live feed; the app hot-reloads it."),
        updated_at="2026-02-07T00:00:00+09:00",
        races={},
        referendums={},
    ))

    # ---------------- election.json
    n_smd = sum(len(r["candidates"]) for r in smd_races)
    n_pr = sum(len(c["list"]) for r in pr_races for c in r["candidates"])
    n_pr_only = sum(1 for r in pr_races for c in r["candidates"] for e in c["list"] if not e.get("smd"))
    write_json(os.path.join(out_dir, "election.json"), OrderedDict(
        schema="jpy_election.election/v1",
        election=ELECTION,
        dissolution="2026-01-23",
        offices=[
            OrderedDict(name_ja="小選挙区", name_en="Single-member districts", name_zh="小選舉區", kind="smd",
                        seats=289, candidates=n_smd, headline=True),
            OrderedDict(name_ja="比例代表", name_en="Proportional representation (11 blocs)",
                        name_zh="比例代表（11區塊）", kind="pr", seats=176, candidates=n_pr,
                        headline=True),
        ],
        totals=OrderedDict(seats=465, majority=233, supermajority=310, candidates=n_smd + n_pr_only),
        blocs=[OrderedDict(id=b["id"], name_ja=b["ja"], name_en=b["en"], name_zh=b["zh"] + "區塊",
                           seats=b["seats"], prefs=b["prefs"]) for b in jp_names.BLOC_BY_ID.values()],
        referendums=referendums,
        estimates=("Eligible voters per district = prefecture total x district share of Japanese population "
                   "(2020 census, 2022 redistricting); per unit = district total split by valid votes."),
        sources=[
            OrderedDict(title="総務省: 衆議院議員総選挙・最高裁判所裁判官国民審査 速報結果/結果調 (令和8年2月8日執行)",
                        url="https://www.soumu.go.jp/senkyo/senkyo_s/data/shugiin51/index.html"),
            OrderedDict(title="総務省: 市区町村別得票数 (第51回)",
                        url="https://www.soumu.go.jp/senkyo/senkyo_s/data/shugiin51/shikuchouson.html"),
            OrderedDict(title="総務省: 市区町村別得票数 (第50回, 2024) — incumbents",
                        url="https://www.soumu.go.jp/senkyo/senkyo_s/data/shugiin50/shikuchouson.html"),
            OrderedDict(title="西澤明 (地域・交通データ研究所): 衆議院小選挙区 2022 区割りポリゴン",
                        url="https://gtfs-gis.jp/senkyoku2022/"),
            OrderedDict(title="SmartNews SMRI japan-topography (国土数値情報 行政区域データ N03, 国土交通省)",
                        url="https://github.com/smartnews-smri/japan-topography"),
            OrderedDict(title="総務省: 全国地方公共団体コード",
                        url="https://www.soumu.go.jp/denshijiti/code.html"),
            OrderedDict(title="Wikipedia: 2026 Japanese general election",
                        url="https://en.wikipedia.org/wiki/2026_Japanese_general_election"),
        ],
    ))
    print(f"  wrote parties.json ({len(parties)}), candidates.json ({len(smd_races)} SMD + {len(pr_races)} PR), "
          f"election.json, results.json, results_preelection.json")
